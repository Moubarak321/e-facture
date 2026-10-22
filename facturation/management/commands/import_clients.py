"""Importe des clients depuis un export Excel de contacts Odoo.

Utilisation (depuis le dossier du projet, environnement virtuel activé) :

    python manage.py import_clients "C:\\chemin\\Contact_odoo_clients.xlsx" --dry-run   # simulation
    python manage.py import_clients "C:\\chemin\\Contact_odoo_clients.xlsx"             # import réel

Correspondance des colonnes :
    Nom d'affichage (ou Nom complet)  ->  Raison sociale
    N° TVA                            ->  IFU (accepté seulement s'il compte 13 chiffres)
    Téléphone + Email                 ->  Contact  (« téléphone · email »)
    Ville                             ->  Adresse

Le script peut être relancé sans risque : un client déjà présent (même IFU, sinon même
raison sociale) n'est jamais dupliqué ni écrasé ; seuls ses champs vides sont complétés.
"""
import csv
import re
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from facturation.models import Client

NAME_COLS = ("Nom d'affichage", "Nom complet")
IFU_COL, PHONE_COL, EMAIL_COL, CITY_COL = "N° TVA", "Téléphone", "Email", "Ville"
MAX_NAME, MAX_CONTACT = 200, 100


def clean(value):
    """Texte sur une ligne, sans espaces superflus (ni espaces insécables)."""
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\u00a0", " ")).strip()


def normalize_ifu(raw):
    """Retourne (ifu, problème). L'IFU béninois compte exactement 13 chiffres."""
    raw = clean(raw)
    if not raw:
        return "", None
    compact = re.sub(r"[\s.\-]", "", raw)  # « 320247833 8352 » -> « 3202478338352 »
    if re.fullmatch(r"\d{13}", compact):
        return compact, None
    return "", f"N° TVA non conforme à un IFU de 13 chiffres : « {raw} »"


def read_rows(path, sheet=None):
    """Lit le fichier Excel et retourne (lignes, anomalies)."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [clean(h) for h in next(rows, [])]

    def col(*names):
        for n in names:
            if n in header:
                return header.index(n)
        return None

    i_name, i_ifu = col(*NAME_COLS), col(IFU_COL)
    i_phone, i_mail, i_city = col(PHONE_COL), col(EMAIL_COL), col(CITY_COL)
    if i_name is None:
        raise CommandError(f"Colonne du nom introuvable (attendu : {' ou '.join(NAME_COLS)}).")

    def get(row, idx):
        return clean(row[idx]) if idx is not None and idx < len(row) else ""

    data, issues = [], []
    for line, row in enumerate(rows, start=2):
        if row is None or not any(c not in (None, "") for c in row):
            continue
        name = get(row, i_name)
        if not name:
            issues.append((line, "", "Ligne ignorée : raison sociale vide", ""))
            continue
        ifu, problem = normalize_ifu(get(row, i_ifu))
        if problem:
            issues.append((line, name, problem, get(row, i_ifu)))
        contact = " · ".join(p for p in (get(row, i_phone), get(row, i_mail)) if p)
        if len(name) > MAX_NAME:
            issues.append((line, name, f"Nom tronqué à {MAX_NAME} caractères", name))
            name = name[:MAX_NAME]
        if len(contact) > MAX_CONTACT:
            issues.append((line, name, f"Contact tronqué à {MAX_CONTACT} caractères", contact))
            contact = contact[:MAX_CONTACT]
        data.append({"line": line, "name": name, "ifu": ifu,
                     "contact": contact, "address": get(row, i_city)})
    return data, issues


class Command(BaseCommand):
    help = "Importe des clients depuis un fichier Excel (export de contacts Odoo)."

    def add_arguments(self, parser):
        parser.add_argument("fichier", help="Chemin du fichier .xlsx")
        parser.add_argument("--feuille", default=None, help="Nom de la feuille (1re par défaut)")
        parser.add_argument("--dry-run", action="store_true",
                            help="Simule l'import sans rien écrire dans la base")
        parser.add_argument("--rapport", default="rapport_import_clients.csv",
                            help="Fichier CSV listant les anomalies")

    def handle(self, *args, **opts):
        path = Path(opts["fichier"])
        if not path.exists():
            raise CommandError(f"Fichier introuvable : {path}")

        data, issues = read_rows(path, opts["feuille"])
        created = updated = unchanged = 0

        with transaction.atomic():
            for r in data:
                existing = Client.objects.filter(ifu=r["ifu"]).first() if r["ifu"] else None
                if existing is None:
                    existing = Client.objects.filter(name__iexact=r["name"]).first()

                if existing is None:
                    Client.objects.create(name=r["name"], ifu=r["ifu"],
                                          contact=r["contact"], address=r["address"])
                    created += 1
                    continue

                changed = False
                for field in ("ifu", "contact", "address"):  # complète seulement les champs vides
                    if r[field] and not getattr(existing, field):
                        setattr(existing, field, r[field])
                        changed = True
                if changed:
                    existing.save()
                    updated += 1
                else:
                    unchanged += 1
            if opts["dry_run"]:
                transaction.set_rollback(True)

        if issues:
            with open(opts["rapport"], "w", newline="", encoding="utf-8-sig") as f:
                w = csv.writer(f, delimiter=";")
                w.writerow(["Ligne Excel", "Client", "Anomalie", "Valeur d'origine"])
                w.writerows(issues)

        mode = "SIMULATION (rien n'a été enregistré)" if opts["dry_run"] else "IMPORT TERMINÉ"
        self.stdout.write(self.style.SUCCESS(mode))
        self.stdout.write(f"  Lignes lues       : {len(data)}")
        self.stdout.write(f"  Clients créés     : {created}")
        self.stdout.write(f"  Clients complétés : {updated}")
        self.stdout.write(f"  Déjà présents     : {unchanged}")
        self.stdout.write(f"  Anomalies         : {len(issues)}"
                          + (f"  (détail dans {opts['rapport']})" if issues else ""))
